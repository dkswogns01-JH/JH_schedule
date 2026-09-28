import calendar
from datetime import date, time
import io
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
import pandas as pd
import streamlit as st

# 웹 페이지 설정
st.set_page_config(
    page_title="근무표 자동 생성기", page_icon="📅", layout="centered"
)

st.title("📅 NCC파견 근무표 생성 시스템")
st.markdown("정보를 입력하고 버튼을 누르세요.")

# 세션 상태 초기화
if "submitted" not in st.session_state:
  st.session_state.submitted = False

# 입력 폼 생성
with st.form("schedule_form"):
  input_name = st.text_input("작업자명 (예: 安材訓)", value="安材訓")
  input_company = st.text_input(
      "회사명 (예: 株式会社 PLANIT)", value="株式会社 PLANIT"
  )

  col1, col2 = st.columns(2)
  with col1:
    input_year = st.number_input("년도", value=2026, step=1)
  with col2:
    input_month = st.number_input(
        "월", value=9, min_value=1, max_value=12, step=1
    )

  submitted_btn = st.form_submit_button("📋 근무표 생성 및 수정하기")

if submitted_btn:
  st.session_state.submitted = True
  st.session_state.input_name = input_name
  st.session_state.input_company = input_company
  st.session_state.input_year = int(input_year)
  st.session_state.input_month = int(input_month)

if st.session_state.submitted:
  cur_year = st.session_state.input_year
  cur_month = st.session_state.input_month
  cur_name = st.session_state.input_name
  cur_company = st.session_state.input_company

  st.markdown("---")
  st.subheader(f"📋 {cur_year}년 {cur_month}월 근무표 미리보기 및 수정")
  st.info(
      "아래 표의 **'出社/休日'** 컬럼에서 **休日**로 선택된 날짜에 시간이"
      " 입력되어 있으면 다운로드가 차단됩니다."
  )

  # 해당 월의 날짜 데이터를 Pandas DataFrame으로 생성
  _, last_day = calendar.monthrange(cur_year, cur_month)
  preview_data = []

  for day in range(1, last_day + 1):
    curr_date = date(cur_year, cur_month, day)
    wd_num = curr_date.weekday()
    is_off = wd_num >= 5

    weekday_kr = ["월", "화", "수", "목", "금", "토", "일"][wd_num]

    preview_data.append({
        "일": day,
        "曜": weekday_kr,
        "근무구분": "休日" if is_off else "出社",
        "시작시간": "" if is_off else "09:00",
        "종료시간": "" if is_off else "18:00",
        "휴게시간": "" if is_off else "01:00",
        "出張先": "",
        "備考": "",
    })

  df_preview = pd.DataFrame(preview_data)


  # 웹 화면 표(DataFrame) 스타일 함수
  def highlight_holiday(val):
    if val == "休日":
      return "color: #FF0000; font-weight: bold;"
    elif val == "休出":
      return "color: #FF8C00; font-weight: bold;"
    return "color: #000000;"


  # Streamlit 데이터 에디터
  edited_df = st.data_editor(
      df_preview.style.map(highlight_holiday, subset=["근무구분"]),
      num_rows="fixed",
      use_container_width=True,
      key="schedule_grid_editor",
      column_config={
          "근무구분": st.column_config.SelectboxColumn(
              "出社/休日",
              help="근무 상태를 선택하세요",
              options=["出社", "休出", "休日"],
              required=True,
          ),
          "일": st.column_config.NumberColumn("日", disabled=True),
          "曜": st.column_config.TextColumn("曜", disabled=True),
      },
  )

  if st.button("📥 수정된 내용으로 엑셀 파일 생성 및 다운로드"):
    # 1. 유효성 검사 (휴일인데 시간이 입력되어 있는지 체크)
    has_error = False
    error_message = ""

    for idx, row_data in edited_df.iterrows():
      day = int(row_data["일"])
      val_c = str(row_data["근무구분"])
      start_val = str(row_data["시작시간"]).strip()
      end_val = str(row_data["종료시간"]).strip()

      if val_c == "休日" and (
          (start_val and start_val != "None" and start_val != "nan")
          or (end_val and end_val != "None" and end_val != "nan")
      ):
        has_error = True
        error_message = f"[실패] {day}일은 '休日(휴일)'로 설정되어 있으나 시간이 입력되어 있습니다. 시간을 비우거나 근무구분을 변경해주세요."
        break

    # 2. 에러가 있을 경우 실패 메시지 출력 및 차단
    if has_error:
      st.error(f"❌ 다운로드 실패! {error_message}")
    else:
      # 3. 정상일 때 엑셀 파일 생성
      wb = openpyxl.Workbook()
      ws = wb.active
      ws.title = "勤　　休　　表"

      ws.views.sheetView[0].showGridLines = True
      ms_pgothic_font = Font(name="MS PGothic", size=11)

      # 1. 타이틀 (Row 1)
      ws.merge_cells("A1:I1")
      ws["A1"] = "勤　　休　　表"
      ws["A1"].font = Font(size=16)
      ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

      # 스타일 정의 (선 종류: thin, medium 및 thin_border 추가)
      yellow_fill = PatternFill(
          start_color="FFFFCC", end_color="FFFFCC", fill_type="solid"
      )
      gray_fill = PatternFill(
          start_color="BFBFBF", end_color="BFBFBF", fill_type="solid"
      )
      orange_fill = PatternFill(
          start_color="FFC000", end_color="FFC000", fill_type="solid"
      )

      thin_side = Side(style="thin", color="000000")
      medium_side = Side(style="medium", color="000000")
      thin_border = Border(
          left=thin_side, right=thin_side, top=thin_side, bottom=thin_side
      )

      # 2. 년도 / 월 (Row 2)
      ws.merge_cells("A2:B2")
      ws["A2"] = cur_year
      ws["A2"].number_format = '0"年"'
      ws["C2"] = cur_month
      ws["C2"].number_format = '0"月"'

      ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
      ws["C2"].alignment = Alignment(horizontal="right", vertical="center")
      ws["A2"].font = Font(size=12, bold=True)
      ws["C2"].font = Font(size=12, bold=True)
      ws["A2"].fill = yellow_fill
      ws["B2"].fill = yellow_fill
      ws["C2"].fill = yellow_fill

      # 3. 회사명 및 작업자명 (Row 4)
      ws.merge_cells("A4:B4")
      ws["A4"] = "会社名："
      ws["A4"].alignment = Alignment(horizontal="right", vertical="bottom")
      ws["B4"].alignment = Alignment(horizontal="right", vertical="bottom")

      ws["C4"] = cur_company
      ws.merge_cells("C4:G4")
      ws["C4"].fill = yellow_fill
      for col_idx in range(3, 6):
        ws.cell(row=4, column=col_idx).fill = yellow_fill
      ws["C4"].alignment = Alignment(horizontal="left", vertical="bottom")

      ws["H4"] = "作業者名："
      ws["I4"] = cur_name

      ws["I4"].fill = yellow_fill
      ws["I4"].alignment = Alignment(horizontal="left", vertical="bottom")
      ws["H4"].alignment = Alignment(horizontal="right", vertical="bottom")

      # 4. 테이블 헤더 (Row 6 & 7)
      ws.merge_cells("A6:A7")
      ws["A6"] = "日"
      ws.merge_cells("B6:B7")
      ws["B6"] = "曜"
      ws.merge_cells("C6:C7")
      ws["C6"] = "出社\n休日"

      ws.merge_cells("D6:E6")
      ws["D6"] = "作業"
      ws["D7"] = "開始時間"
      ws["E7"] = "終了時刻"

      ws.merge_cells("F6:F7")
      ws["F6"] = "休憩\n時間"
      ws.merge_cells("G6:G7")
      ws["G6"] = "実作業\n時間"
      ws.merge_cells("H6:H7")
      ws["H6"] = "出張先"
      ws.merge_cells("I6:I7")
      ws["I6"] = "備考"

      for row in range(6, 8):
        for col in range(1, 10):
          cell = ws.cell(row=row, column=col)
          cell.fill = gray_fill
          cell.alignment = Alignment(
              horizontal="center", vertical="center", wrap_text=True
          )

      # 헤더 영역(6~7행, A~I열) 박스 테두리 적용
      for r in range(6, 8):
        for c in range(1, 10):
          cell = ws.cell(row=r, column=c)
          t_top = medium_side if r == 6 else thin_side
          t_bottom = medium_side if r == 7 else thin_side
          t_left = medium_side if c == 1 else thin_side
          t_right = medium_side if c == 9 else thin_side
          cell.border = Border(
              top=t_top, bottom=t_bottom, left=t_left, right=t_right
          )

      # 5. 데이터 행 생성 (Row 8 ~ 말일)
      start_row = 8
      end_row = start_row + last_day - 1

      dv = DataValidation(
          type="list", formula1='"出社,休出,休日"', allow_blank=True
      )
      ws.add_data_validation(dv)
      dv.add(f"C{start_row}:C{end_row}")

      for idx, row_data in edited_df.iterrows():
        r = start_row + idx
        day = int(row_data["일"])
        curr_date = date(cur_year, cur_month, day)
        wd_num = curr_date.weekday()

        cell_a = ws.cell(row=r, column=1, value=f"=DATE($A$2,$C$2,{day})")
        cell_a.number_format = "d"

        ws.cell(row=r, column=2, value=f'=TEXT(A{r},"aaa")')

        val_c = str(row_data["근무구분"])
        cell_c = ws.cell(row=r, column=3, value=val_c)

        if val_c == "休日":
          cell_c.font = Font(color="FF0000")
          ws.cell(row=r, column=4, value="")
          ws.cell(row=r, column=5, value="")
          ws.cell(row=r, column=6, value="")
        else:
          try:
            start_t = (
                time.fromisoformat(str(row_data["시작시간"]))
                if pd.notna(row_data["시작시간"])
                and str(row_data["시작시간"]) != ""
                and str(row_data["시작시간"]) != "nan"
                else None
            )
          except:
            start_t = None

          try:
            end_t = (
                time.fromisoformat(str(row_data["종료시간"]))
                if pd.notna(row_data["종료시간"])
                and str(row_data["종료시간"]) != ""
                and str(row_data["종료시간"]) != "nan"
                else None
            )
          except:
            end_t = None

          try:
            break_t = (
                time.fromisoformat(str(row_data["휴게시간"]))
                if pd.notna(row_data["휴게시간"])
                and str(row_data["휴게시간"]) != ""
                and str(row_data["휴게시간"]) != "nan"
                else None
            )
          except:
            break_t = None

          if start_t:
            cell_d = ws.cell(row=r, column=4, value=start_t)
            cell_d.number_format = "HH:mm"
            cell_d.fill = yellow_fill
          else:
            ws.cell(row=r, column=4, value="")

          if end_t:
            cell_e = ws.cell(row=r, column=5, value=end_t)
            cell_e.number_format = "HH:mm"
            cell_e.fill = yellow_fill
          else:
            ws.cell(row=r, column=5, value="")

          if break_t:
            cell_f = ws.cell(row=r, column=6, value=break_t)
            cell_f.number_format = "HH:mm"
            cell_f.fill = yellow_fill
          else:
            ws.cell(row=r, column=6, value="")

          cell_c.fill = yellow_fill

        if val_c != "休日":
          cell_g = ws.cell(row=r, column=7, value=f"=MOD(E{r}-D{r},1)-F{r}")
        else:
          cell_g = ws.cell(row=r, column=7, value="")
        cell_g.number_format = "HH:mm"

        ws.cell(
            row=r,
            column=8,
            value=str(row_data["出張先"]) if pd.notna(row_data["出張先"]) else "",
        )
        ws.cell(
            row=r,
            column=9,
            value=str(row_data["備考"]) if pd.notna(row_data["備考"]) else "",
        )

        # 데이터 영역(8행~말일, A~I열) 박스 테두리 적용
        for c_idx in range(1, 10):
          cell = ws.cell(row=r, column=c_idx)
          t_top = medium_side if r == start_row else thin_side
          t_bottom = medium_side if r == end_row else thin_side
          t_left = medium_side if c_idx == 1 else thin_side
          t_right = medium_side if c_idx == 9 else thin_side
          cell.border = Border(
              top=t_top, bottom=t_bottom, left=t_left, right=t_right
          )

          if c_idx in [1, 2, 3]:
            cell.alignment = Alignment(horizontal="center", vertical="center")
          else:
            cell.alignment = Alignment(horizontal="right", vertical="center")

          if c_idx == 3 and val_c == "休日":
            cell.font = Font(color="FF0000")

      # 6. 맨 하단 합계 행 추가
      total_row = end_row + 1

      cell_total_g = ws.cell(
          row=total_row, column=7, value=f"=SUM(G{start_row}:G{end_row})"
      )
      cell_total_g.number_format = "[h]:mm"
      cell_total_g.fill = orange_fill
      cell_total_g.alignment = Alignment(horizontal="center", vertical="center")

      cell_total_h = ws.cell(
          row=total_row, column=8, value=f"=G{total_row}*24"
      )
      cell_total_h.number_format = '0.00"時間"'
      cell_total_h.fill = orange_fill
      cell_total_h.alignment = Alignment(horizontal="center", vertical="center")

      # 하단 합계 행(G~I열) 박스 테두리 적용
      for c_idx in range(1, 10):
        cell = ws.cell(row=total_row, column=c_idx)
        if c_idx in [7, 8, 9]:
          t_top = medium_side
          t_bottom = medium_side
          t_left = medium_side if c_idx == 7 else thin_side
          t_right = medium_side if c_idx == 9 else thin_side
          cell.border = Border(
              top=t_top, bottom=t_bottom, left=t_left, right=t_right
          )
          if c_idx == 9:
            cell.alignment = Alignment(horizontal="center", vertical="center")
        else:
          cell.border = Border()

      # 7. 우측 하단 도장칸 추가
      stamp_start_row = total_row + 2
      stamp_title = ws.cell(
          row=stamp_start_row, column=9, value="NCC承認印"
      )
      stamp_title.font = Font(bold=True, size=10)
      stamp_title.alignment = Alignment(horizontal="center", vertical="center")

      for r_offset in range(1, 7):
        ws.cell(row=stamp_start_row + r_offset, column=9)

      for r_idx in range(stamp_start_row, stamp_start_row + 7):
        ws.cell(row=r_idx, column=9).border = thin_border

      for row in range(6, ws.max_row + 1):
        ws.row_dimensions[row].height = 18

      for row in ws.iter_rows():
        for cell in row:
          has_red_color = (
              cell.font
              and cell.font.color
              and cell.font.color.rgb
              and "FF0000" in str(cell.font.color.rgb)
          )
          if has_red_color:
            cell.font = Font(name="MS PGothic", size=10, color="FF0000")
          else:
            cell.font = Font(name="MS PGothic", size=10, color="000000")

      #반영 작게되서 일단 밑에 형식
    #   ws.column_dimensions["A"].width = 5.00
    #   ws.column_dimensions["B"].width = 5.00
    #   ws.column_dimensions["C"].width = 5.14
    #   ws.column_dimensions["D"].width = 8.14
    #   ws.column_dimensions["E"].width = 8.00
    #   ws.column_dimensions["F"].width = 8.00
    #   ws.column_dimensions["G"].width = 8.00
    #   ws.column_dimensions["H"].width = 14.29
    #   ws.column_dimensions["I"].width = 21.00

      ws.column_dimensions["A"].width = 5.8
      ws.column_dimensions["B"].width = 5.8
      ws.column_dimensions["C"].width = 8.84
      ws.column_dimensions["D"].width = 8.84
      ws.column_dimensions["E"].width = 8.80
      ws.column_dimensions["F"].width = 8.80
      ws.column_dimensions["G"].width = 8.80
      ws.column_dimensions["H"].width = 14.5
      ws.column_dimensions["I"].width = 23.10


      ws.row_dimensions[1].height = 41.25
      ws.row_dimensions[2].height = 21.75
      ws.row_dimensions[3].height = 7.5
      ws.row_dimensions[4].height = 22.5
      ws.row_dimensions[5].height = 9

      ws.cell(row=1, column=1).font = Font(
          name="MS PGothic", size=18, color="000000"
      )

      ws.cell(row=2, column=1).font = Font(
          name="MS PGothic", size=12, bold=True, color="000000"
      )
      ws.cell(row=2, column=2).font = Font(
          name="MS PGothic", size=12, bold=True, color="000000"
      )
      ws.cell(row=2, column=3).font = Font(
          name="MS PGothic", size=12, bold=True, color="000000"
      )

      ws.cell(row=4, column=1).font = Font(
          name="MS PGothic", size=11, color="000000"
      )
      ws.cell(row=4, column=2).font = Font(
          name="MS PGothic", size=11, color="000000"
      )
      ws.cell(row=4, column=3).font = Font(
          name="MS PGothic", size=11, color="000000"
      )
      ws.cell(row=4, column=4).font = Font(
          name="MS PGothic", size=11, color="000000"
      )
      ws.cell(row=4, column=8).font = Font(
          name="MS PGothic", size=11, color="000000"
      )
      ws.cell(row=4, column=9).font = Font(
          name="MS PGothic", size=11, color="000000"
      )



      code_row = stamp_start_row + 7
      code_cell = ws.cell(row=code_row, column=9, value="NCC-HC4-260701")
      code_cell.font = Font(size=8)
      code_cell.alignment = Alignment(horizontal="right", vertical="bottom")

      # 메모리에 엑셀 파일 저장
      buffer = io.BytesIO()
      wb.save(buffer)
      buffer.seek(0)

      filename = f"work_schedule_{cur_name}_{cur_year}_{cur_month}월.xlsx"

      st.success("엑셀 파일이 성공적으로 생성되었습니다!")
      st.download_button(
          label="📥 파일 다운로드 받기",
          data=buffer,
          file_name=filename,
          mime=(
              "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
          ),
      )