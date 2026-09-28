import calendar
from datetime import date, time
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation


def create_work_schedule(name, company, year, month):
  wb = openpyxl.Workbook()
  ws = wb.active
  ws.title = "勤　　休　　表"

  # 엑셀 격자선(Gridlines) 표시 설정
  ws.views.sheetView[0].showGridLines = True

  # 사용할 폰트 정의 (이름과 크기 지정)
  ms_pgothic_font = Font(name="MS PGothic", size=11)

  # 1. 타이틀 (Row 1)
  ws.merge_cells("A1:I1")
  ws["A1"] = "勤　　休　　表"
  ws["A1"].font = Font(size=16, bold=True)
  ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

  # 스타일 정의
  yellow_fill = PatternFill(
      start_color="FFFFCC", end_color="FFFFCC", fill_type="solid"
  )
  gray_fill = PatternFill(
      start_color="BFBFBF", end_color="BFBFBF", fill_type="solid"
  )
  thin_border = Border(
      left=Side(style="thin", color="000000"),
      right=Side(style="thin", color="000000"),
      top=Side(style="thin", color="000000"),
      bottom=Side(style="thin", color="000000"),
  )

  # 2. 년도 / 월 (Row 2) - A2:B2 병합, C2는 월 단독
  ws.merge_cells("A2:B2")
  ws["A2"] = year
  ws["A2"].number_format = '0"年"'

  ws["C2"] = month
  ws["C2"].number_format = '0"月"'

  ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
  ws["C2"].alignment = Alignment(horizontal="right", vertical="center")
  ws["A2"].font = Font(size=12, bold=True)
  ws["C2"].font = Font(size=12, bold=True)
  ws["A2"].fill = yellow_fill
  ws["B2"].fill = yellow_fill
  ws["C2"].fill = yellow_fill

  # 3. 회사명 및 작업자명 (Row 4)
  ws.merge_cells("A4:B4")
  ws["A4"] = "会社名："
  ws["A4"].font = Font(bold=True)
  ws["A4"].alignment = Alignment(horizontal="right", vertical="bottom")
  ws["B4"].alignment = Alignment(horizontal="right", vertical="bottom")

  ws["C4"] = company
  ws.merge_cells("C4:G4")
  ws["C4"].fill = yellow_fill
  for col_idx in range(3, 6):
    ws.cell(row=4, column=col_idx).fill = yellow_fill
  ws["C4"].alignment = Alignment(horizontal="left", vertical="bottom")

  ws["H4"] = "作業者名："
  ws["I4"] = name

  ws["I4"].fill = yellow_fill
  ws["H4"].font = Font(bold=True)
  ws["I4"].alignment = Alignment(horizontal="left", vertical="bottom")
  ws["H4"].alignment = Alignment(horizontal="right", vertical="bottom")

  # 4. 테이블 헤더 (Row 6 & 7)
  ws.merge_cells("A6:A7")
  ws["A6"] = "日"

  ws.merge_cells("B6:B7")
  ws["B6"] = "曜"

  ws.merge_cells("C6:C7")
  ws["C6"] = "出社\n休日"

  ws.merge_cells("D6:E6")
  ws["D6"] = "作業"
  ws["D7"] = "開始時間"
  ws["E7"] = "終了時刻"

  ws.merge_cells("F6:F7")
  ws["F6"] = "休憩\n時間"

  ws.merge_cells("G6:G7")
  ws["G6"] = "実作業\n時間"

  ws.merge_cells("H6:H7")
  ws["H6"] = "出張先"

  ws.merge_cells("I6:I7")
  ws["I6"] = "備考"

  for row in range(6, 8):
    for col in range(1, 10):
      cell = ws.cell(row=row, column=col)
      cell.fill = gray_fill
      cell.font = Font(bold=True)
      cell.alignment = Alignment(
          horizontal="center", vertical="center", wrap_text=True
      )
      cell.border = thin_border

  # 5. 데이터 행 생성 (Row 8 ~ 말일)
  _, last_day = calendar.monthrange(year, month)
  start_row = 8
  end_row = start_row + last_day - 1

  # C열 드롭다운 설정
  dv = DataValidation(type="list", formula1='"出社,休出,休日"', allow_blank=True)
  ws.add_data_validation(dv)
  dv.add(f"C{start_row}:C{end_row}")

  for day in range(1, last_day + 1):
    r = start_row + day - 1
    curr_date = date(year, month, day)
    wd_num = curr_date.weekday()
    is_off = wd_num >= 5

    # A열: 날짜 수식
    cell_a = ws.cell(row=r, column=1, value=f"=DATE($A$2,$C$2,{day})")
    cell_a.number_format = "d"

    # B열: 요일 수식
    ws.cell(row=r, column=2, value=f'=TEXT(A{r},"aaa")')

    # C열: 출근/휴일 구분 (휴일인 경우 빨간 글씨 적용)
    val_c = "休日" if is_off else "出社"
    cell_c = ws.cell(row=r, column=3, value=val_c)

    if is_off:
      cell_c.font = Font(color="FF0000", bold=True)

    if is_off:
      ws.cell(row=r, column=4, value="")
      ws.cell(row=r, column=5, value="")
      ws.cell(row=r, column=6, value="")
    else:
      cell_d = ws.cell(row=r, column=4, value=time(9, 0))
      cell_d.number_format = "HH:mm"

      cell_e = ws.cell(row=r, column=5, value=time(18, 0))
      cell_e.number_format = "HH:mm"

      cell_f = ws.cell(row=r, column=6, value=time(1, 0))
      cell_f.number_format = "HH:mm"
      
      cell_c.fill = yellow_fill
      cell_d.fill = yellow_fill
      cell_e.fill = yellow_fill
      cell_f.fill = yellow_fill

    # G열: 실작업 시간 수식
    cell_g = ws.cell(row=r, column=7, value=f"=MOD(E{r}-D{r},1)-F{r}")
    cell_g.number_format = "HH:mm"

    ws.cell(row=r, column=8, value="")
    ws.cell(row=r, column=9, value="")

    for c_idx in range(1, 10):
      cell = ws.cell(row=r, column=c_idx)
      cell.border = thin_border
      if c_idx in [1, 2, 3]:
        cell.alignment = Alignment(horizontal="center", vertical="center")
      else: 
        cell.alignment = Alignment(horizontal="right", vertical="center")
      if c_idx == 2 and is_off:
        cell.font = Font(color="FF0000", bold=True)
      elif c_idx == 3 and is_off:
        cell.font = Font(color="FF0000", bold=True)

  # 6. 맨 하단 합계 행 추가 (말일 바로 다음 행)
  total_row = end_row + 1

  # G열 합계 수식 및 서식
  cell_total_g = ws.cell(
      row=total_row, column=7, value=f"=SUM(G{start_row}:G{end_row})"
  )
  cell_total_g.number_format = "[h]:mm"
  cell_total_g.font = Font(bold=True)
  cell_total_g.fill = yellow_fill
  cell_total_g.alignment = Alignment(horizontal="center", vertical="center")
  cell_total_g.border = thin_border

  # H열 소수점 시간 환산 수식 및 사용자지정 서식
  cell_total_h = ws.cell(row=total_row, column=8, value=f"=G{total_row}*24")
  cell_total_h.number_format = '0.00"時間"'
  cell_total_h.font = Font(bold=True)
  cell_total_h.fill = yellow_fill
  cell_total_h.alignment = Alignment(horizontal="center", vertical="center")
  cell_total_h.border = thin_border

  # A~F열은 합계 행 테두리를 제거하고, I열은 테두리 유지
  for c_idx in range(1, 10):
    cell = ws.cell(row=total_row, column=c_idx)
    if c_idx in [7, 8, 9]:
      cell.border = thin_border
      if c_idx == 9:
        cell.alignment = Alignment(horizontal="center", vertical="center")
    else:
      cell.border = Border()  # 테두리 없음

  # 7. 우측 하단 도장칸 추가 (6칸으로 확장)
  stamp_start_row = total_row + 2

  # I열 단독 'NCC 承認印' 타이틀
  stamp_title = ws.cell(row=stamp_start_row, column=9, value="NCC承認印")
  stamp_title.font = Font(bold=True, size=10)
  stamp_title.alignment = Alignment(horizontal="center", vertical="center")

  # 도장 서명용 빈 칸들 (I열 단독으로 6개 행)
  for r_offset in range(1, 7):
    r = stamp_start_row + r_offset
    # 빈 칸 셀 객체 확보 (필요시 테두리 적용을 위해 변수로 받아도 좋습니다)
    ws.cell(row=r, column=9)

  # 도장칸 전체 테두리 박스 설정 (I열만 해당하므로 c_idx를 9로 고정)
  for r_idx in range(stamp_start_row, stamp_start_row + 7):
    ws.cell(row=r_idx, column=9).border = thin_border


  # 1. 전체 행 높이 일괄 적용
  for row in range(6, ws.max_row + 1):
    ws.row_dimensions[row].height = 18

  # 2. [마지막 루프] 전체 기본 폰트 일괄 적용 (빨간색 요일 컬럼 보호하기)
  for row in ws.iter_rows():
    for cell in row:
      # 셀에 폰트와 컬러 정보가 있고, 색상 코드에 'FF0000'(빨강)이 포함되어 있다면?
      has_red_color = (
          cell.font
          and cell.font.color
          and cell.font.color.rgb
          and "FF0000" in str(cell.font.color.rgb)
      )

      if has_red_color:
        # 빨간색 글씨(요일 등)
        cell.font = Font(
            name="MS PGothic",
            size=10,
            color="FF0000",
        )
      else:
        # 나머지는 일반 기본 폰트 적용
        cell.font = Font(name="MS PGothic", size=10, color="000000")
   # 열 너비 크기 조절
    ws.column_dimensions["A"].width = 5.00
    ws.column_dimensions["B"].width = 5.00
    ws.column_dimensions["C"].width = 5.14
    ws.column_dimensions["D"].width = 8.14
    ws.column_dimensions["E"].width = 8.00
    ws.column_dimensions["F"].width = 8.00
    ws.column_dimensions["G"].width = 8.00
    ws.column_dimensions["H"].width = 14.29
    ws.column_dimensions["I"].width = 21.00

    # 행 높이 조절
    ws.row_dimensions[1].height = 41.25
    ws.row_dimensions[2].height = 21.75
    ws.row_dimensions[3].height = 7.5
    ws.row_dimensions[4].height = 22.5
    ws.row_dimensions[5].height = 9

    # 헤더쪽 폰트 조절
    cell = ws.cell(row=1, column=1)  # 1행 1열
    cell.font = Font(name="MS PGothic", size=18, bold=True, color="000000")

    cell2 = ws.cell(row=2, column=1)  # 2행 1열
    cell2.font = Font(name="MS PGothic", size=12, bold=True, color="000000")

    cell2 = ws.cell(row=2, column=3)  # 2행 3열
    cell2.font = Font(name="MS PGothic", size=12, bold=True, color="000000")

    # 도장칸 맨 하단 코드 텍스트 (I열 단독)
    code_row = stamp_start_row + 7
    code_cell = ws.cell(row=code_row, column=9, value="NCC-HC4-260701")
    code_cell.font = Font(size=8)
    code_cell.alignment = Alignment(horizontal="right", vertical="bottom")

  # 파일 저장
  filename = f"work_schedule_{name}_{year}_{month}월.xlsx"
  wb.save(filename)
  print(
      f"\n[완료] {name}님의 {year}년 {month}월 근무표(도장칸 6칸 확장)가"
      f" 생성되었습니다: {filename}"
  )


if __name__ == "__main__":
  input_name = input("작업자명을 입력하세요 (예: 安材訓): ")
  input_company = input("회사명을 입력하세요 (예: 株式会社 PLANIT): ")
  input_year = int(input("년도를 입력하세요 (예: 2026): "))
  input_month = int(input("월을 입력하세요 (예: 9): "))

  create_work_schedule(input_name, input_company, input_year, input_month)